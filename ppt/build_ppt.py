"""Build the defense PPTX from ppt/slides.md.

Run with:
    uv run --no-project --python 3.13.13 --with python-pptx python ppt/build_ppt.py
"""

from __future__ import annotations

import math
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
ASSETS = HERE / "assets"
OUT = HERE / "out"
OUT.mkdir(exist_ok=True)

FONT = "Microsoft YaHei"
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.62)
TITLE_TOP = Inches(0.34)
TITLE_H = Inches(0.86)
BODY_TOP = Inches(1.36)
BODY_BOTTOM = Inches(7.02)

PRIMARY = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = RGBColor(0x2E, 0x8B, 0x8B)
TEXT = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x5A, 0x66, 0x72)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xE8, 0xEE, 0xF4)
LIGHT2 = RGBColor(0xF4, 0xF7, 0xFA)
WARN = RGBColor(0xB0, 0x60, 0x00)


def In(v: float) -> Emu:
    return Inches(v)


# --------------------------------------------------------------------------
# parsing
# --------------------------------------------------------------------------

def parse_meta(raw: str) -> dict:
    meta: dict = {}
    for line in raw.splitlines():
        line = line.strip()
        if not line or ":" not in line:
            continue
        key, _, value = line.partition(":")
        meta[key.strip()] = value.strip()
    return meta


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for i, line in enumerate(lines):
        if i == 1 and set(line.replace("|", "").replace(" ", "")) <= set("-:"):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


def parse_blocks(text: str) -> list[dict]:
    blocks: list[dict] = []
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.lstrip().startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            blocks.append({"type": "table", "rows": parse_table(table_lines)})
            continue
        if line.lstrip().startswith("- "):
            items = []
            while i < len(lines) and lines[i].lstrip().startswith("- "):
                items.append(lines[i].lstrip()[2:].strip())
                i += 1
            blocks.append({"type": "ul", "items": items})
            continue
        m = re.match(r"^\s*(\d+)\.\s+(.*)$", line)
        if m:
            items = []
            while i < len(lines):
                mm = re.match(r"^\s*(\d+)\.\s+(.*)$", lines[i])
                if not mm:
                    break
                items.append(mm.group(2).strip())
                i += 1
            blocks.append({"type": "ol", "items": items})
            continue
        blocks.append({"type": "p", "text": line.strip()})
        i += 1
    return blocks


def parse_slides(text: str) -> tuple[dict, list[dict]]:
    deck_m = re.search(r"<!--\s*deck(.*?)-->", text, re.S)
    deck = parse_meta(deck_m.group(1)) if deck_m else {}

    slides = []
    pattern = re.compile(r"<!--\s*slide(.*?)-->(.*?)(?=<!--\s*slide|$)", re.S)
    for m in pattern.finditer(text):
        meta = parse_meta(m.group(1))
        body = m.group(2)
        title_m = re.search(r"^##\s+(.*)$", body, re.M)
        title = title_m.group(1).strip() if title_m else ""
        sections: dict[str, str] = {}
        parts = re.split(r"^###\s+(.*)$", body, flags=re.M)
        for j in range(1, len(parts), 2):
            sections[parts[j].strip()] = parts[j + 1]
        slides.append(
            {
                "meta": meta,
                "title": title,
                "blocks": parse_blocks(sections.get("正文", "")),
                "notes": sections.get("备注", "").strip(),
            }
        )
    return deck, slides


# --------------------------------------------------------------------------
# low level helpers
# --------------------------------------------------------------------------

def set_run(run, size: float, bold: bool = False, color: RGBColor = TEXT) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", FONT)


def add_label(slide, left, top, width, height, text, size=14, bold=False,
              color=TEXT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.3
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run(run, size, bold, color)
    return tb


def add_rect(slide, left, top, width, height, text, size=11, fill=LIGHT,
             line=PRIMARY, color=TEXT, bold=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    try:
        shorter = min(width, height) / 914400.0
        shp.adjustments[0] = max(0.0, min(0.5, 0.07 / shorter)) if shorter > 0 else 0.06
    except (IndexError, ValueError):
        pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    shp.line.width = Pt(1.0)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = In(0.05)
    tf.margin_right = In(0.05)
    tf.margin_top = In(0.02)
    tf.margin_bottom = In(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for idx, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run(run, size, bold, color)
    return shp


def add_arrow(slide, x1, y1, x2, y2, color=PRIMARY, width=1.4):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(x1), int(y1), int(x2), int(y2))
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return conn


def display_len(text: str) -> float:
    n = 0.0
    for ch in text:
        n += 1.0 if ord(ch) > 0x2E7F else 0.55
    return n


def est_lines(text: str, width, size: float) -> int:
    win = width / 914400.0
    per_line = max(1.0, win / (size / 72.0))
    return max(1, math.ceil(display_len(text) / per_line))


def est_height(text: str, width, size: float) -> int:
    lines = est_lines(text, width, size)
    return int(lines * (size / 72.0) * 1.5 * 914400)


# --------------------------------------------------------------------------
# block rendering
# --------------------------------------------------------------------------

def render_table(slide, rows, left, top, width, size=12):
    nrows = len(rows)
    ncols = max(len(r) for r in rows)
    row_h = In(0.34)
    table_h = int(row_h * nrows)
    gf = slide.shapes.add_table(nrows, ncols, int(left), int(top), int(width), table_h)
    table = gf.table
    maxlens = []
    for c in range(ncols):
        maxlens.append(max(display_len(r[c]) if c < len(r) else 0 for r in rows))
    total = sum(maxlens) or 1
    for c in range(ncols):
        table.columns[c].width = int(width * maxlens[c] / total)
    for r in range(nrows):
        table.rows[r].height = int(row_h)
        for c in range(ncols):
            cell = table.cell(r, c)
            cell.text = rows[r][c] if c < len(rows[r]) else ""
            cell.margin_left = In(0.06)
            cell.margin_right = In(0.06)
            cell.margin_top = In(0.01)
            cell.margin_bottom = In(0.01)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = PRIMARY
            else:
                cell.fill.fore_color.rgb = LIGHT2 if r % 2 else WHITE
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER if r == 0 else PP_ALIGN.LEFT
                p.line_spacing = 1.15
                p.space_after = Pt(0)
                for run in p.runs:
                    set_run(run, size, bold=(r == 0), color=WHITE if r == 0 else TEXT)
    return table_h


def text_block_height(blocks, width, size):
    h = 0.0
    for b in blocks:
        if b["type"] == "p":
            h += est_lines(b["text"], width, size) * (size / 72.0) * 1.5 + 0.06
        elif b["type"] in ("ul", "ol"):
            for item in b["items"]:
                h += est_lines("• " + item, width, size) * (size / 72.0) * 1.5 + 0.035
    return int(h * 914400)


def render_text(slide, blocks, left, top, width, size):
    height = text_block_height(blocks, width, size)
    tb = slide.shapes.add_textbox(int(left), int(top), int(width), height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    first = True
    for b in blocks:
        if b["type"] == "p":
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.line_spacing = 1.3
            p.space_after = Pt(5)
            run = p.add_run()
            run.text = b["text"]
            set_run(run, size)
        elif b["type"] in ("ul", "ol"):
            for idx, item in enumerate(b["items"]):
                p = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                p.line_spacing = 1.3
                p.space_after = Pt(2)
                marker = "• " if b["type"] == "ul" else f"{idx + 1}. "
                run = p.add_run()
                run.text = marker + item
                set_run(run, size)
    return height


def render_blocks(slide, blocks, left, top, width, bottom, size=15, table_size=11):
    text_blocks = [b for b in blocks if b["type"] in ("p", "ul", "ol")]
    table_blocks = [b for b in blocks if b["type"] == "table"]
    y = top
    if text_blocks:
        h = render_text(slide, text_blocks, left, top, width, size)
        y = top + h + In(0.12)
    for tb in table_blocks:
        h = render_table(slide, tb["rows"], left, y, width, size=table_size)
        y += h + In(0.12)
    return y



# --------------------------------------------------------------------------
# figures
# --------------------------------------------------------------------------

def add_image(slide, path, L, T, W, H):
    p = Path(path)
    if not p.is_absolute():
        p = ROOT / path
    if p.exists():
        pic = slide.shapes.add_picture(str(p), int(L), int(T))
        ratio = pic.width / pic.height
        box_ratio = W / H
        if ratio >= box_ratio:
            pic.width = int(W)
            pic.height = int(W / ratio)
        else:
            pic.height = int(H)
            pic.width = int(H * ratio)
        pic.left = int(L + (W - pic.width) / 2)
        pic.top = int(T + (H - pic.height) / 2)
    else:
        add_rect(slide, L, T, W, H, f"图片待补充\n{path}", 11,
                 fill=LIGHT2, line=WARN, color=WARN)


def add_figure(slide, fig, L, T, W, H):
    if fig["kind"] == "image":
        add_image(slide, fig["value"], L, T, W, H)
    else:
        path = ASSETS / f"diagram-{fig['value']}.png"
        if path.exists():
            add_image(slide, str(path), L, T, W, H)
        else:
            add_rect(slide, L, T, W, H, f"图形待补充\n{fig['value']}", 11,
                     fill=LIGHT2, line=WARN, color=WARN)


def collect_figures(meta: dict) -> list[dict]:
    figs = []
    if meta.get("diagram"):
        for d in meta["diagram"].split(";"):
            d = d.strip()
            if d:
                figs.append({"kind": "diagram", "value": d})
    if meta.get("素材"):
        for s in meta["素材"].split(";"):
            s = s.strip()
            if s:
                figs.append({"kind": "image", "value": s})
    return figs


# --------------------------------------------------------------------------
# slide builders
# --------------------------------------------------------------------------

def add_title(slide, title, page):
    add_label(slide, MARGIN, TITLE_TOP, SLIDE_W - 2 * MARGIN, TITLE_H,
              title, 27, True, PRIMARY, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(MARGIN), int(TITLE_TOP + TITLE_H),
                                 int(In(1.1)), int(In(0.05)))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    bar.shadow.inherit = False


def add_footer(slide, deck_title, page):
    add_label(slide, MARGIN, In(7.08), SLIDE_W - 2 * MARGIN, In(0.3),
              deck_title, 9, False, MUTED)
    add_label(slide, SLIDE_W - MARGIN - In(1.0), In(7.08), In(1.0), In(0.3),
              str(page), 9, False, MUTED, PP_ALIGN.RIGHT)


def build_cover(slide, title, blocks, deck):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, int(SLIDE_W), int(In(2.4)))
    bg.fill.solid()
    bg.fill.fore_color.rgb = PRIMARY
    bg.line.fill.background()
    bg.shadow.inherit = False
    body = blocks
    main = body[0]["text"] if body and body[0]["type"] == "p" else title
    add_label(slide, MARGIN, In(0.85), SLIDE_W - 2 * MARGIN, In(1.2),
              main, 26, True, WHITE, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    y = In(2.85)
    for idx, b in enumerate(body[1:]):
        if b["type"] == "p":
            size = 18 if idx == 0 else 15
            color = ACCENT if idx == 0 else TEXT
            add_label(slide, MARGIN, y, SLIDE_W - 2 * MARGIN, In(0.5),
                      b["text"], size, False, color)
            y += In(0.6)


def build_section(slide, title, blocks):
    add_label(slide, MARGIN, In(2.7), SLIDE_W - 2 * MARGIN, In(1.0),
              title, 34, True, PRIMARY, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(MARGIN), int(In(3.85)),
                                 int(In(2.0)), int(In(0.07)))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    bar.shadow.inherit = False


def build_end(slide, title, blocks):
    add_label(slide, MARGIN, In(1.0), SLIDE_W - 2 * MARGIN, In(0.9),
              title, 30, True, PRIMARY, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    render_blocks(slide, blocks, MARGIN, In(2.1), SLIDE_W - 2 * MARGIN, In(6.4), size=17)


def build_content(slide, title, blocks, meta, page, deck):
    add_title(slide, title, page)
    figs = collect_figures(meta)
    pos = meta.get("figure_pos", "right")
    body_left = MARGIN
    body_w = SLIDE_W - 2 * MARGIN
    body_bottom = BODY_BOTTOM
    size = 15
    if figs and pos == "right":
        body_w = In(6.9)
        size = 15
        render_blocks(slide, blocks, body_left, BODY_TOP, body_w, BODY_BOTTOM, size=size,
                      table_size=10)
        fig_left = MARGIN + body_w + In(0.3)
        fig_w = SLIDE_W - MARGIN - fig_left
        fig_top = BODY_TOP + In(0.1)
        fig_h = BODY_BOTTOM - fig_top - In(0.1)
        add_figure(slide, figs[0], fig_left, fig_top, fig_w, fig_h)
    elif figs and pos == "bottom":
        render_blocks(slide, blocks, body_left, BODY_TOP, body_w, In(3.55), size=14,
                      table_size=10)
        fig_top = In(3.65)
        fig_h = BODY_BOTTOM - fig_top
        add_figure(slide, figs[0], MARGIN, fig_top, SLIDE_W - 2 * MARGIN, fig_h)
    elif figs and pos == "dual-bottom":
        render_blocks(slide, blocks, body_left, BODY_TOP, body_w, In(3.55), size=14,
                      table_size=10)
        fig_top = In(3.65)
        fig_h = BODY_BOTTOM - fig_top
        gap = In(0.3)
        fw = (SLIDE_W - 2 * MARGIN - gap) / 2
        add_figure(slide, figs[0], MARGIN, fig_top, fw, fig_h)
        if len(figs) > 1:
            add_figure(slide, figs[1], MARGIN + fw + gap, fig_top, fw, fig_h)
    elif figs and pos == "right-bottom":
        text_w = In(6.5)
        render_blocks(slide, blocks, body_left, BODY_TOP, text_w, In(4.25), size=14,
                      table_size=10)
        fig_left = MARGIN + text_w + In(0.3)
        fig_w = SLIDE_W - MARGIN - fig_left
        add_figure(slide, figs[0], fig_left, BODY_TOP, fig_w, In(2.95))
        if len(figs) > 1:
            b_top = In(4.4)
            b_h = BODY_BOTTOM - b_top
            add_figure(slide, figs[1], MARGIN, b_top, SLIDE_W - 2 * MARGIN, b_h)
    else:
        render_blocks(slide, blocks, body_left, BODY_TOP, body_w, BODY_BOTTOM, size=16,
                      table_size=11)
    add_footer(slide, deck.get("title", ""), page)


def build_demo(slide, title, blocks, meta, page, deck):
    build_content(slide, title, blocks, meta, page, deck)


# --------------------------------------------------------------------------
# main
# --------------------------------------------------------------------------

def main():
    text = (HERE / "slides.md").read_text(encoding="utf-8")
    deck, slides = parse_slides(text)

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    for page, item in enumerate(slides, start=1):
        slide = prs.slides.add_slide(blank)
        layout = item["meta"].get("layout", "content")
        title = item["title"]
        blocks = item["blocks"]
        if layout == "cover":
            build_cover(slide, title, blocks, deck)
        elif layout == "section":
            build_section(slide, title, blocks)
        elif layout == "end":
            build_end(slide, title, blocks)
        else:
            build_content(slide, title, blocks, item["meta"], page, deck)
        if item["notes"]:
            slide.notes_slide.notes_text_frame.text = item["notes"]

    out = OUT / "课程设计_验收答辩.pptx"
    try:
        prs.save(str(out))
    except PermissionError:
        out = OUT / "课程设计_验收答辩_新.pptx"
        prs.save(str(out))
        print("原文件被占用，已另存为新文件")
    print(f"slides: {len(slides)}")
    print(f"saved: {out}")


if __name__ == "__main__":
    main()
